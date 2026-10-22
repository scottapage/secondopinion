"""The sales dossier: an XLSX of every service with its porting details and recommended 2degrees plan,
plus the customer's bill PDFs, emailed to the sales team."""
import io
import os
import re
import time
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import analysis, db, learning

NAVY, BLUE, LIGHT, ORANGE = "003660", "009ED8", "E0F3FA", "F1600D"
MAX_ATTACH_BYTES = 20 * 1024 * 1024


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _money(cell):
    cell.number_format = '"$"#,##0.00'


def build_xlsx(email):
    """Returns (bytes, analysis) for the customer's whole estate."""
    a = analysis.build(email=email)
    user = db.q1("SELECT * FROM users WHERE email=?", (email,)) or {}
    sw = a.get("switch", {"lines": [], "services": []})
    wb = Workbook()

    # --- Services (the sheet sales asked for)
    ws = wb.active
    ws.title = "Services"
    cols = ["Provider", "LSP Account Number", "LSP Account Name", "MSISDN", "Username", "Current Plan",
            "Current Spend (excl GST/mth)", "Extras (avg/mth)", "Recommended 2degrees Plan", "Role", "Future Price",
            "Monthly Saving", "Basis", "Port-ready"]
    ws.append(cols)
    _style_header(ws, 1, len(cols))
    for l in sw["lines"]:
        ready = "Already 2degrees" if l["already"] else ("Yes - needs user name" if not l["user"] or
                                                           l["user"] == l.get("account_name") else "Yes")
        ws.append([l["carrier"], l["account"], l.get("account_name"), l["display"], l["user"] or "", l["plan"],
                   l["current"], l["extras"], l["rec"] + (f" ({l['alt']} if usage fits)" if l.get("alt") else ""),
                   l["role"], l["price"], l["saving"], l["why"], ready])
        r = ws.max_row
        for c in (7, 8, 11, 12):
            _money(ws.cell(row=r, column=c))
        if l["saving"] < 0:
            ws.cell(row=r, column=12).font = Font(color="B5470A")
    for s in sw["services"]:
        ws.append([s["carrier"], "", "", "", s["address"] or "", f"{s['type']}: {s['plan'] or ''}", s["current"], 0,
                   s["rec"], "", s["price"], s["saving"], s["why"], "n/a (broadband)"])
        r = ws.max_row
        for c in (7, 8, 11, 12):
            _money(ws.cell(row=r, column=c))
    ws.append([])
    ws.append(["", "", "", "", "", "TOTAL (plans + add-ons)", sw["current_month"], sw["extras_month"], "", "",
               sw["future_month"], sw["saving_month"], f"Right-sized scenario: {sw['optimised_month']:.2f}/mth"])
    r = ws.max_row
    for c in (7, 8, 11, 12):
        _money(ws.cell(row=r, column=c))
        ws.cell(row=r, column=c).font = Font(bold=True)
    ws.cell(row=r, column=6).font = Font(bold=True)
    _autosize(ws, [11, 18, 26, 15, 22, 30, 16, 12, 34, 12, 12, 12, 44, 20])
    ws.freeze_panes = "A2"

    # --- Summary
    ws2 = wb.create_sheet("Summary")
    k = a.get("kpi", {})
    score = a.get("score", {})
    rows = [
        ("Customer", user.get("company") or ""), ("Contact", user.get("name") or ""), ("Email", email),
        ("Providers", ", ".join(f"{p['name']} ({p['lines']} lines, {p['accounts']} account(s))" for p in a.get("providers", []))),
        ("Accounts", "; ".join(f"{x['carrier']} {x['number']} - {x['name']}" for x in a.get("accounts", []))),
        ("Bills analysed", f"{k.get('bills', 0)} ({', '.join(a.get('month_labels', []))})"),
        ("Current mobile run-rate (excl GST)", k.get("monthly")), ("Yearly run-rate", k.get("annual")),
        ("Broadband & other services (avg/mth)", k.get("other_services")),
        ("Device repayments (per month)", k.get("devices_month")),
        ("Hardware balance remaining", sw.get("hardware_balance")),
        ("Mobile health score", f"{score.get('score', '')} / 100 - {score.get('grade', '')}"),
        ("Like-for-like on 2degrees (per month)", sw.get("future_month")),
        ("Like-for-like saving (per year)", sw.get("saving_year")),
        ("Right-sized on 2degrees (per month)", sw.get("optimised_month")),
        ("Right-sized saving (per year)", sw.get("optimised_saving_year")),
        ("Findings", "; ".join(f"{i['title']}: {i['headline']}" for i in a.get("insights", []))),
    ]
    ws2.append(["Item", "Value"])
    _style_header(ws2, 1, 2)
    for key, val in rows:
        ws2.append([key, val])
        if isinstance(val, (int, float)) and "score" not in key.lower():
            _money(ws2.cell(row=ws2.max_row, column=2))
    for n in learning.notes_for(email):
        ws2.append([f"{'Did you know' if n.get('kind') == 'fact' else 'Expert note'}: {n['title']}", n["body"]])
    if sw.get("roaming"):
        r = sw["roaming"]
        ws2.append(["Roaming now vs 2degrees (per month)", f"{r['now_month']:.2f} now vs ~{r['future_month']:.2f} on 2degrees"
                    f" ({r['days_other']} days outside Australia at ${r['rate']:.0f}/day, {r['days_au']} days in Australia free)"])
    for b in sw.get("benefits", []):
        ws2.append([f"2degrees {b['category']}: {b['title']}", f"{b['value']} - {b['detail']} ({b['conditions']})".strip(" -")])
    _autosize(ws2, [38, 110])
    for row in ws2.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # --- Porting checklist
    ws3 = wb.create_sheet("Porting")
    pcols = ["MSISDN", "Losing provider", "LSP account number", "LSP account name", "User name", "Delivery address",
             "eSIM?", "Status"]
    ws3.append(pcols)
    _style_header(ws3, 1, len(pcols))
    for l in sw["lines"]:
        if l["already"]:
            continue
        name_ok = bool(l["user"]) and l["user"] != l.get("account_name")
        ws3.append([l["display"], l["carrier"], l["account"], l.get("account_name"),
                    l["user"] if name_ok else "NEEDED", "NEEDED (or eSIM)", "",
                    "Ready to port once user name confirmed" if not name_ok else "Ready to port"])
    _autosize(ws3, [15, 14, 18, 28, 22, 22, 8, 36])

    # --- Devices
    ws4 = wb.create_sheet("Devices")
    ws4.append(["Line", "Device", "Monthly", "Payments left", "Balance remaining"])
    _style_header(ws4, 1, 5)
    for d in a.get("devices", []):
        ws4.append([d["number"], d["description"], d["monthly"], d["months_left"], d["balance"]])
        _money(ws4.cell(row=ws4.max_row, column=3))
        _money(ws4.cell(row=ws4.max_row, column=5))
    _autosize(ws4, [15, 50, 12, 14, 18])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), a


def bill_attachments(email):
    """(name, bytes, mime) for each distinct bill PDF, newest first, within the size cap."""
    rows = db.q("SELECT * FROM invoices WHERE owner_email=? AND status IN ('done','duplicate') AND path<>'-'"
                " ORDER BY COALESCE(period_month,'') DESC, id DESC", (email,))
    out, seen, total = [], set(), 0
    for r in rows:
        if r["sha256"] in seen or not os.path.exists(r["path"]):
            continue
        seen.add(r["sha256"])
        data = open(r["path"], "rb").read()
        if total + len(data) > MAX_ATTACH_BYTES:
            break
        total += len(data)
        name = f"{r['carrier'] or 'bill'}-{r['account_number'] or ''}-{r['period_month'] or ''}-{r['id']}.pdf"
        out.append((re.sub(r"[^\w.-]", "_", name), data, "application/pdf"))
    return out, len(rows) - len(out)


def filename(email):
    stamp = datetime.now(ZoneInfo("Pacific/Auckland")).strftime("%Y%m%d")
    return f"SecondOpinion-dossier-{re.sub(r'[^a-z0-9]+', '-', email.lower())}-{stamp}.xlsx"


def send(email, req=None, to=None):
    """Email the dossier (XLSX + bill PDFs) to sales. Returns the analysis used."""
    from . import mailer
    from .tpl import render
    xlsx, a = build_xlsx(email)
    pdfs, skipped = bill_attachments(email)
    attachments = [(filename(email), xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")] + pdfs
    to = to or db.get_setting("sales_email")
    user = db.q1("SELECT * FROM users WHERE email=?", (email,))
    offer = db.q1("SELECT * FROM offers WHERE id=?", (req["offer_id"],)) if req and req.get("offer_id") else None
    kinds = {"quote": "Quote request", "offer": "Offer claimed", "call": "Call requested"}
    subject = (kinds.get(req["kind"], "Request") if req else "Sales dossier") + ": " + \
              ((user or {}).get("company") or email)
    html = render("email/sales.html", req=req or {"kind": "dossier", "email": email, "phone": "", "notes": ""},
                  offer=offer, a=a, user=user, spend_chart=None, skipped=skipped, dossier=True)
    mailer.send(to, subject, html, reply_to=email, attachments=attachments)
    db.log_event("dossier_sent", email, to=to, bills=len(pdfs))
    return a
