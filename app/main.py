import asyncio
import json
import os
import re
import secrets
import subprocess
import threading
import time
from pathlib import Path

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from starlette.exceptions import HTTPException as StarletteHTTPException
from starlette.middleware.sessions import SessionMiddleware

from . import (analysis, catalogue, chat, db, dossier, evidence, example, extract, ingest, learning, mailer, meter, notify,
               security, share, visitors)
from .config import (BASE_URL, CARRIERS, CHART_DIR, DATA_DIR, MAX_UPLOAD_MB, SECRET_KEY, SECRET_SETTINGS, SETUP_TOKEN,
                     SUPERADMIN_EMAIL, carrier_label)
from .tpl import render

db.init()
app = FastAPI(docs_url=None, redoc_url=None, openapi_url=None)
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, max_age=60 * 60 * 24 * 60, same_site="lax",
                   https_only=os.environ.get("BASE_URL", "").startswith("https"))
app.mount("/static", StaticFiles(directory=Path(__file__).parent / "static"), name="static")

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
try:
    VERSION = subprocess.run(["git", "rev-parse", "--short", "HEAD"], capture_output=True, text=True,
                             cwd=Path(__file__).parent).stdout.strip() or "dev"
    BRANCH = subprocess.run(["git", "rev-parse", "--abbrev-ref", "HEAD"], capture_output=True, text=True,
                            cwd=Path(__file__).parent).stdout.strip() or "main"
except Exception:
    VERSION, BRANCH = "dev", "main"
DEPLOY_BRANCH_FILE = DATA_DIR / "deploy-branch"
SITE_VERSIONS = {"main": "V2 - current site", "v1": "V1 - first release (frozen)"}


def deploy_branch():
    try:
        return DEPLOY_BRANCH_FILE.read_text().strip() or "main"
    except OSError:
        return "main"


def set_deploy_branch(branch):
    """The auto-deploy script (deploy/update.sh) deploys whatever branch this file names, within a minute."""
    if branch not in SITE_VERSIONS:
        raise HTTPException(400, "Unknown site version")
    DEPLOY_BRANCH_FILE.write_text(branch + "\n")
    db.log_event("site_version", None, None, branch=branch)


BOT_RE = re.compile(r"bot|crawl|spider|slurp|preview|monitor|curl|python|headless|wget|httpx", re.I)
SKIP_PREFIXES = ("/static", "/api/", "/c/", "/healthz", "/setup", "/invoice/", "/favicon")


@app.middleware("http")
async def track(request: Request, call_next):
    """Record unhandled errors in the events table (page views come from the browser beacon, /api/pv)."""
    path = request.url.path
    try:
        response = await call_next(request)
    except Exception as e:
        db.log_event("error", None, None, path=path, error=f"{type(e).__name__}: {e}"[:300])
        raise
    if response.status_code >= 500:
        db.log_event("error", None, None, path=path, status=response.status_code)
    return response


# ---------- helpers ----------

def client_ip(request):
    """Caddy sets X-Forwarded-For and uvicorn runs with --proxy-headers, so request.client is the visitor."""
    xff = request.headers.get("x-forwarded-for", "")
    return (xff.split(",")[0].strip() if xff else (request.client.host if request.client else "")) or ""


def sid(request):
    s = request.session.get("sid")
    if not s:
        s = request.session["sid"] = secrets.token_urlsafe(16)
    return s


def current_email(request):
    return request.session.get("email")


def is_admin(request):
    return current_email(request) == SUPERADMIN_EMAIL


def is_locked(request):
    """Front-end password on, and this visitor hasn't entered it (the superadmin is never locked out)."""
    if request.url.path in ("/login", "/forgot", "/verify", "/resend", "/set-password") or request.url.path.startswith("/admin"):
        return False  # the superadmin must always be able to log in and switch the lock off
    return db.get_setting("site_lock") == "on" and not request.session.get("unlocked") and not is_admin(request)


def require_unlocked(request):
    if is_locked(request):
        raise HTTPException(403, "This site is password protected.")


def page(request, name, status_code=200, **ctx):
    ctx.update(request=request, user_email=current_email(request), is_admin=is_admin(request),
               mail_address=db.get_setting("mail_address"), version=VERSION, locked=is_locked(request),
               lock_error=request.session.pop("lock_error", None))
    return HTMLResponse(render(name, **ctx), status_code)


def require_admin(request):
    if not is_admin(request):
        raise HTTPException(404)


def login(request, email):
    email = email.lower()
    request.session["email"] = email
    ingest.claim_session(request.session.get("sid"), email)
    db.x("UPDATE users SET last_login=? WHERE email=?", (time.time(), email))


def safe_next(n):
    return n if n and n.startswith("/") and not n.startswith("//") else "/dashboard"


def background(fn, *args):
    threading.Thread(target=fn, args=args, daemon=True).start()


# ---------- public pages ----------

@app.get("/healthz")
def healthz():
    return {"ok": True, "version": VERSION}


@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    sid(request)
    return page(request, "index.html", meter=meter.get(), carriers=CARRIERS)


@app.get("/api/meter")
def api_meter():
    return JSONResponse(meter.get(), headers={"Cache-Control": "public, max-age=300"})


# ---------- shareable cards ----------

@app.post("/api/share")
async def api_share(request: Request):
    require_unlocked(request)
    body = await request.json()
    kind, key = str(body.get("kind") or "")[:20], str(body.get("key") or "")[:30]
    email = current_email(request)
    token = share.create(kind, key, owner_email=email, session_id=None if email else sid(request))
    if not token:
        raise HTTPException(400, "Nothing to share yet")
    card = share.get(token)
    return {"url": share.url(token), "title": card["title"], "big": card["big"], "body": card["body"],
            "image": f"{BASE_URL}/s/{token}.png"}


@app.get("/s/{token}.png")
def share_image(token: str):
    card = share.get(re.sub(r"[^\w-]", "", token))
    if not card:
        raise HTTPException(404)
    return FileResponse(share.image(card), media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})


@app.get("/s/{token}", response_class=HTMLResponse)
def share_page(request: Request, token: str):
    card = share.get(re.sub(r"[^\w-]", "", token))
    if not card:
        raise HTTPException(404)
    db.x("UPDATE share_cards SET views=views+1 WHERE id=?", (card["id"],))
    db.log_event("share_viewed", None, sid(request), card=card["kind"])
    return page(request, "share.html", card=card)


@app.post("/unlock")
def unlock(request: Request, password: str = Form(...), next: str = Form("/")):
    if secrets.compare_digest(password.strip(), db.get_setting("site_password") or ""):
        request.session["unlocked"] = True
        db.log_event("site_unlocked", current_email(request), sid(request), ip=client_ip(request))
    else:
        request.session["lock_error"] = "That password didn't match."
        db.log_event("site_unlock_failed", None, sid(request), ip=client_ip(request))
    return RedirectResponse(safe_next(next) if next != "/dashboard" or current_email(request) else "/", 303)


@app.post("/api/upload")
async def upload(request: Request, files: list[UploadFile] = File(...)):
    require_unlocked(request)
    email = current_email(request)
    session_id = sid(request)
    sub_id, token = ingest.new_submission("web", email=email, session_id=session_id)
    count = 0
    for f in files[:20]:
        data = await f.read()
        if not data or len(data) > MAX_UPLOAD_MB * 1024 * 1024:
            continue
        ingest.add_file(sub_id, data, f.filename, owner_email=email, session_id=session_id)
        count += 1
    if not count:
        raise HTTPException(400, "No usable files")
    db.log_event("upload", email, session_id, files=count, ip=client_ip(request))
    return {"token": token, "redirect": f"/results/{token}"}


def _submission_for(request, token):
    sub = db.q1("SELECT * FROM submissions WHERE token=?", (token,))
    if not sub:
        raise HTTPException(404)
    email = current_email(request)
    if not (is_admin(request) or sub["session_id"] == request.session.get("sid") or (email and sub["email"] == email)):
        raise HTTPException(404)
    return sub


@app.get("/api/submission/{token}")
def submission_status(request: Request, token: str):
    sub = _submission_for(request, token)
    invs = db.q("SELECT id, filename, status, error, carrier, carrier_name, account_name, period_month, total_excl_gst,"
                " line_count FROM invoices WHERE submission_id=? ORDER BY id", (sub["id"],))
    for i in invs:
        i["carrier_label"] = carrier_label(i["carrier"], i.pop("carrier_name"))
    pending = any(i["status"] in ("queued", "processing") for i in invs)
    return {"pending": pending, "invoices": invs}


@app.get("/results/{token}", response_class=HTMLResponse)
def results(request: Request, token: str):
    sub = _submission_for(request, token)
    db.log_event("results_viewed", current_email(request), request.session.get("sid"))
    return page(request, "results.html", token=token, logged_in=bool(current_email(request)), sub=sub)


def _teaser(a):
    """High-level view for people who haven't registered: hide the detail behind the insights."""
    if not a.get("has_data"):
        return a
    for i in a["insights"]:
        i["count"] = len(i["items"])
        i["items"] = []
        i["saving"] = None
    a["kpi"]["savings_month"] = None
    a["kpi"]["savings_year"] = None
    a["devices"] = [{"count": len(a["devices"])}] if a["devices"] else []
    if a.get("switch"):
        a["switch"]["lines"] = []
        a["switch"]["services"] = []
    for l in a["lines"]:
        l["flags"] = []
    a["teaser"] = True
    return a


def _clean(a):
    a = dict(a)
    a["invoices"] = [{k: i[k] for k in ("id", "filename", "status", "error", "carrier", "account_name",
                                        "account_number", "period_month", "total_excl_gst", "line_count",
                                        "invoice_number")} | {"carrier_label": carrier_label(i["carrier"], i.get("carrier_name"))}
                     for i in a.get("invoices", [])]
    return a


@app.get("/api/analysis")
def api_analysis(request: Request, as_email: str = None, account: str = None):
    email = current_email(request)
    if as_email and is_admin(request):
        email = as_email.lower()
    if email:
        a = analysis.build(email=email, account=account or None)
        a["notes"] = learning.notes_for(email)
        return JSONResponse(_clean(a))
    sess = request.session.get("sid")
    a = analysis.build(session_id=sess)
    a["notes"] = [{"title": n["title"]} for n in learning.notes_for(session_id=sess)]
    return JSONResponse(_clean(_teaser(a)))


# ---------- visitor tracking (browser beacon: scanners that don't run the page are not counted) ----------

BEACON_EVENTS = {"upload_intent", "files_selected", "example_cta", "email_form_shown"}


@app.post("/api/pv")
async def page_view_beacon(request: Request):
    body = await request.json()
    ua = request.headers.get("user-agent", "")
    if BOT_RE.search(ua):
        return {"ok": True}
    kind = body.get("kind") or "page_view"
    if kind != "page_view" and kind not in BEACON_EVENTS:
        raise HTTPException(400)
    ref_host = re.sub(r"^https?://([^/]+).*$", r"\1", str(body.get("ref") or ""))[:80]
    if "secondopinion.co.nz" in ref_host or "sslip.io" in ref_host:
        ref_host = ""
    meta = {"js": 1, "path": str(body.get("path") or "")[:120],
            "device": "mobile" if re.search(r"Mobi|Android|iPhone", ua) else "desktop", "ip": client_ip(request)}
    if ref_host:
        meta["ref"] = ref_host
    db.log_event(kind, current_email(request), sid(request), **meta)
    return {"ok": True}


# ---------- public example ----------

@app.get("/example", response_class=HTMLResponse)
def example_page(request: Request):
    db.log_event("example_viewed", current_email(request), sid(request))
    return page(request, "results.html", token="", logged_in=False, sub=None, example=True)


@app.get("/api/example")
def api_example():
    a = analysis.build(email=example.EMAIL)
    a["notes"] = learning.notes_for(example.EMAIL)
    a["example"] = True
    return JSONResponse(_clean(a))


# ---------- "email me this summary" for people who haven't registered ----------

from itsdangerous import BadSignature, URLSafeTimedSerializer  # noqa: E402

CLAIMS = URLSafeTimedSerializer(SECRET_KEY, salt="claim")


@app.post("/api/email-summary")
async def email_summary(request: Request):
    require_unlocked(request)
    body = await request.json()
    email = str(body.get("email") or "").strip().lower()
    if not EMAIL_RE.match(email):
        return JSONResponse({"error": "Please enter a valid email address."}, 400)
    session_id = sid(request)
    last = request.session.get("last_summary_email", 0)
    if time.time() - last < 60:
        return JSONResponse({"error": "We've just sent one - please check your inbox."}, 429)
    request.session["last_summary_email"] = time.time()
    db.x("UPDATE submissions SET contact_email=? WHERE session_id=? AND email IS NULL", (email, session_id))
    token = CLAIMS.dumps({"sid": session_id, "email": email})
    try:
        await asyncio.to_thread(notify.session_summary, email, session_id, f"{BASE_URL}/claim/{token}")
    except Exception as e:
        return JSONResponse({"error": f"We couldn't send the email ({type(e).__name__}). Please try again."}, 500)
    db.log_event("email_captured", email, session_id)
    return {"ok": True}


@app.get("/claim/{token}", response_class=HTMLResponse)
def claim(request: Request, token: str):
    """The link in the summary email. Clicking it proves the person owns the address, so the bills uploaded in
    that browser session are attached to it."""
    try:
        data = CLAIMS.loads(token, max_age=60 * 60 * 24 * 30)
    except BadSignature:
        return _auth(request, "register", error="That link has expired - enter your email to get a new code.")
    email = data["email"]
    ingest.claim_session(data["sid"], email)
    db.log_event("email_verified_by_link", email, data["sid"])
    user = db.q1("SELECT * FROM users WHERE email=?", (email,))
    if user and user["password_hash"]:
        return _auth(request, "login", email=email, next="/dashboard",
                     info="Your bills have been added to your account - log in to see the full analysis.")
    request.session["verified_email"] = email
    return _auth(request, "password", email=email, next="/dashboard", user=user)


@app.get("/c/{name}")
def chart_image(name: str):
    if not re.fullmatch(r"[\w-]+\.png", name):
        raise HTTPException(404)
    p = CHART_DIR / name
    if not p.exists():
        raise HTTPException(404)
    return FileResponse(p, media_type="image/png", headers={"Cache-Control": "public, max-age=31536000"})


def _my_invoice(request, invoice_id):
    """The bill, if this visitor may see it: its owner, the superadmin, or the browser session that uploaded it."""
    inv = db.q1("SELECT * FROM invoices WHERE id=?", (invoice_id,))
    email = current_email(request)
    if not inv or inv["path"] == "-" or not (is_admin(request) or (email and inv["owner_email"] == email) or
                                            (inv["owner_email"] is None and inv["session_id"] == request.session.get("sid"))):
        raise HTTPException(404)
    return inv


@app.get("/bill/{invoice_id}", response_class=HTMLResponse)
def bill_viewer(request: Request, invoice_id: int, q: str = "", label: str = ""):
    """Click any number on the analysis and see it highlighted on the original bill."""
    inv = _my_invoice(request, invoice_id)
    terms = [t.strip() for t in q.split(",") if t.strip()][:6]
    hits = evidence.find(inv, terms)
    found = {}
    for h in hits:
        found.setdefault(h["term"], h["page"])
    # Open on the page where the most terms were found together
    pages_hit = [h["page"] for h in hits]
    first = max(set(pages_hit), key=pages_hit.count) if pages_hit else 0
    db.log_event("evidence_viewed", current_email(request), sid(request), invoice_id=invoice_id, terms=len(terms),
                 found=len(found))
    return page(request, "bill.html", inv=inv, terms=terms, hits=hits, found=found, first=first, label=label[:120],
                pages=evidence.page_count(inv))


@app.get("/bill/{invoice_id}/page/{page_no}.png")
def bill_page_png(request: Request, invoice_id: int, page_no: int):
    inv = _my_invoice(request, invoice_id)
    path = evidence.page_png(inv, page_no)
    if not path:
        raise HTTPException(404)
    return FileResponse(path, media_type="image/png", headers={"Cache-Control": "private, max-age=86400"})


@app.get("/invoice/{invoice_id}.pdf")
def invoice_pdf(request: Request, invoice_id: int):
    inv = _my_invoice(request, invoice_id)
    return FileResponse(inv["path"], media_type="application/pdf", filename=inv["filename"],
                        content_disposition_type="inline")


# ---------- auth: email -> code -> password ----------

def _auth(request, mode, **ctx):
    return page(request, "auth.html", mode=mode, **ctx)


@app.get("/register", response_class=HTMLResponse)
def register_form(request: Request, email: str = "", next: str = ""):
    if current_email(request):
        return RedirectResponse(safe_next(next), 303)
    return _auth(request, "register", email=email, next=next)


@app.post("/register")
def register_submit(request: Request, email: str = Form(...), next: str = Form("")):
    email = email.strip().lower()
    if not EMAIL_RE.match(email):
        return _auth(request, "register", email=email, next=next, error="Please enter a valid email address.")
    user = db.q1("SELECT * FROM users WHERE email=?", (email,))
    if user and user["password_hash"]:
        return _auth(request, "login", email=email, next=next, info="You already have an account - log in below.")
    return _send_code(request, email, "register", next)


def _send_code(request, email, purpose, next=""):
    code = security.issue_code(email, purpose)
    if code:
        try:
            notify.code(email, code, purpose)
        except Exception as e:
            return _auth(request, "register" if purpose == "register" else "forgot", email=email, next=next,
                         error=f"We couldn't send the email ({type(e).__name__}: {str(e)[:160]}). Please try again shortly.")
    request.session["pending_email"] = email
    request.session["pending_purpose"] = purpose
    return _auth(request, "verify", email=email, next=next)


@app.post("/verify")
def verify(request: Request, code: str = Form(...), next: str = Form("")):
    email = request.session.get("pending_email")
    if not email:
        return RedirectResponse("/register", 303)
    if not security.verify_code(email, code):
        return _auth(request, "verify", email=email, next=next, error="That code didn't match or has expired.")
    request.session["verified_email"] = email
    user = db.q1("SELECT * FROM users WHERE email=?", (email,))
    return _auth(request, "password", email=email, next=next, user=user)


@app.post("/resend")
def resend(request: Request, next: str = Form("")):
    email = request.session.get("pending_email")
    if not email:
        return RedirectResponse("/register", 303)
    return _send_code(request, email, request.session.get("pending_purpose", "register"), next)


@app.post("/set-password")
def set_password(request: Request, password: str = Form(...), name: str = Form(""), company: str = Form(""),
                 next: str = Form("")):
    email = request.session.get("verified_email")
    if not email:
        return RedirectResponse("/register", 303)
    if len(password) < 8:
        return _auth(request, "password", email=email, next=next, error="Use at least 8 characters.")
    now = time.time()
    if db.q1("SELECT 1 FROM users WHERE email=?", (email,)):
        db.x("UPDATE users SET password_hash=?, verified_at=COALESCE(verified_at, ?),"
             " name=COALESCE(NULLIF(?,''), name), company=COALESCE(NULLIF(?,''), company) WHERE email=?",
             (security.hash_password(password), now, name.strip(), company.strip(), email))
    else:
        db.x("INSERT INTO users(email, password_hash, name, company, created_at, verified_at) VALUES (?,?,?,?,?,?)",
             (email, security.hash_password(password), name.strip() or None, company.strip() or None, now, now))
        db.log_event("registered", email, request.session.get("sid"), ip=client_ip(request))
        background(notify.admin_alert, f"New account: {email}", [f"{name.strip() or 'no name'} - {company.strip() or 'no company'}"],
                   f"{BASE_URL}/admin/user?email={email}", "Open their analysis", email)
    request.session.pop("verified_email", None)
    request.session.pop("pending_email", None)
    login(request, email)
    return RedirectResponse(safe_next(next), 303)


@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request, next: str = "", email: str = ""):
    if current_email(request):
        return RedirectResponse(safe_next(next), 303)
    return _auth(request, "login", next=next, email=email)


@app.post("/login")
def login_submit(request: Request, email: str = Form(...), password: str = Form(...), next: str = Form("")):
    email = email.strip().lower()
    user = db.q1("SELECT * FROM users WHERE email=?", (email,))
    if not user or not user["password_hash"] or not security.check_password(password, user["password_hash"]):
        return _auth(request, "login", email=email, next=next, error="That email and password don't match.")
    login(request, email)
    return RedirectResponse(safe_next(next), 303)


@app.get("/forgot", response_class=HTMLResponse)
def forgot_form(request: Request, email: str = ""):
    return _auth(request, "forgot", email=email)


@app.post("/forgot")
def forgot_submit(request: Request, email: str = Form(...)):
    email = email.strip().lower()
    if not EMAIL_RE.match(email):
        return _auth(request, "forgot", email=email, error="Please enter a valid email address.")
    return _send_code(request, email, "reset")


@app.get("/logout")
def logout(request: Request):
    request.session.pop("email", None)
    return RedirectResponse("/", 303)


# ---------- signed-in pages ----------

@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(request: Request):
    email = current_email(request)
    if not email:
        return RedirectResponse("/login?next=/dashboard", 303)
    db.log_event("dashboard_viewed", email)
    user = db.q1("SELECT * FROM users WHERE email=?", (email,))
    return page(request, "dashboard.html", user=user, view_email=email, admin_view=False)


@app.post("/api/chat")
async def api_chat(request: Request):
    require_unlocked(request)
    email = current_email(request)
    if not email:
        raise HTTPException(401)
    body = await request.json()
    q = (body.get("q") or "").strip()[:2000]
    if not q:
        raise HTTPException(400)
    try:
        answer = await asyncio.to_thread(chat.ask, email, q)
    except Exception as e:
        answer = f"Sorry - something went wrong answering that ({type(e).__name__}). Please try again."
    return {"answer": answer}


@app.get("/api/chat/history")
def chat_history(request: Request):
    email = current_email(request)
    if not email:
        raise HTTPException(401)
    rows = db.q("SELECT role, content FROM chat_messages WHERE email=? ORDER BY id DESC LIMIT 20", (email,))
    return rows[::-1]


@app.get("/offer", response_class=HTMLResponse)
def offer_page(request: Request):
    email = current_email(request)
    if not email:
        return RedirectResponse("/register?next=/offer", 303)
    offers = db.q("SELECT * FROM offers WHERE active=1 ORDER BY sort, id")
    a = analysis.build(email=email)
    return page(request, "offer.html", offers=offers, published=published_offers(), a=a,
                booking_url=db.get_setting("booking_url"), sent=request.query_params.get("sent"))


def published_offers():
    """Offers as published on 2degrees.nz right now: plan promotions from the catalogue and association discounts
    from the benefits research. Nothing invented."""
    out = []
    for r in db.q("SELECT * FROM plan_catalogue WHERE active=1 AND provider='2degrees' AND promo_months>0"
                  " ORDER BY category, price_excl_gst"):
        out.append({"title": r["promo_text"], "plan": f"{r['plan_name']} - {analysis.fmt_money(r['price_excl_gst'], 2)}/mth excl GST",
                    "badge": "Business " + r["category"], "detail": (r["conditions"] or "")[:220], "url": r["source_url"],
                    "key": f"promo:{r['plan_name']}"})
    for b in db.q("SELECT * FROM benefits WHERE active=1 AND provider='2degrees' AND category='discount' ORDER BY id"):
        if re.search(r"chamber|member|association|partner", (b["title"] or "").lower()):
            out.append({"title": b["title"], "plan": b["value"] or "", "badge": "Members", "detail": (b["detail"] or "")[:220],
                        "url": b["source_url"], "key": f"benefit:{b['title']}"})
    return out[:8]


@app.post("/offer")
def offer_submit(request: Request, kind: str = Form(...), offer_id: str = Form(""), phone: str = Form(""),
                 notes: str = Form("")):
    email = current_email(request)
    if not email:
        return RedirectResponse("/register?next=/offer", 303)
    if kind not in ("quote", "offer", "call"):
        raise HTTPException(400)
    oid = int(offer_id) if offer_id.isdigit() else None
    rid = db.x("INSERT INTO offer_requests(email, kind, offer_id, phone, notes, created_at) VALUES (?,?,?,?,?,?)",
               (email, kind, oid, phone[:40], notes[:2000], time.time()))
    db.log_event("offer_request", email, kind=kind, offer_id=oid)
    background(notify.offer_request, rid)
    background(notify.admin_alert, f"Offer request ({kind}) from {email}", [phone, notes[:300]],
               f"{BASE_URL}/admin/requests", "Open offer requests", email)
    return RedirectResponse(f"/offer?sent={kind}", 303)


# ---------- first-run setup (secrets are entered here, never committed) ----------

@app.get("/setup", response_class=HTMLResponse)
def setup_form(request: Request, token: str = ""):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    return page(request, "setup.html", token=token, configured={k: bool(db.get_setting(k)) for k in SECRET_SETTINGS},
                mail_user=db.get_setting("mail_user"), saved=request.query_params.get("saved"))


@app.post("/setup")
def setup_submit(request: Request, token: str = Form(...), anthropic_api_key: str = Form(""),
                 mail_password: str = Form(""), mail_user: str = Form(""), admin_password: str = Form("")):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    if len(admin_password) >= 8:
        now = time.time()
        db.x("INSERT INTO users(email, password_hash, created_at, verified_at) VALUES (?,?,?,?)"
             " ON CONFLICT(email) DO UPDATE SET password_hash=excluded.password_hash",
             (SUPERADMIN_EMAIL, security.hash_password(admin_password), now, now))
    if anthropic_api_key.strip():
        db.set_setting("anthropic_api_key", anthropic_api_key.strip())
    if mail_password.strip():
        db.set_setting("mail_password", mail_password.replace(" ", "").strip())
    if mail_user.strip():
        db.set_setting("mail_user", mail_user.strip().lower())
    return RedirectResponse(f"/setup?token={token}&saved=1", 303)


# ---------- superadmin ----------

@app.get("/admin", response_class=HTMLResponse)
def admin_home(request: Request):
    require_admin(request)
    recent = db.q("SELECT i.*, s.source FROM invoices i LEFT JOIN submissions s ON s.id=i.submission_id"
                  " ORDER BY i.id DESC LIMIT 12")
    return page(request, "admin/home.html", s=analysis.admin_stats(), recent=recent, tab="home", live=activity(24))


@app.get("/admin/invoices", response_class=HTMLResponse)
def admin_invoices(request: Request, status: str = "", carrier: str = "", qs: str = ""):
    require_admin(request)
    where, args = ["1=1"], []
    if status:
        where.append("i.status=?")
        args.append(status)
    if carrier:
        where.append("i.carrier=?")
        args.append(carrier)
    if qs:
        where.append("(i.owner_email LIKE ? OR i.account_name LIKE ? OR i.account_number LIKE ? OR i.filename LIKE ?)")
        args += [f"%{qs}%"] * 4
    rows = db.q(f"SELECT i.*, s.source FROM invoices i LEFT JOIN submissions s ON s.id=i.submission_id"
                f" WHERE {' AND '.join(where)} ORDER BY i.id DESC LIMIT 300", args)
    return page(request, "admin/invoices.html", rows=rows, status=status, carrier=carrier, qs=qs, tab="invoices",
                carriers=CARRIERS)


@app.get("/admin/invoice/{invoice_id}", response_class=HTMLResponse)
def admin_invoice(request: Request, invoice_id: int):
    require_admin(request)
    inv = db.q1("SELECT * FROM invoices WHERE id=?", (invoice_id,))
    if not inv:
        raise HTTPException(404)
    lines = db.q("SELECT * FROM lines WHERE invoice_id=?", (invoice_id,))
    usage = db.q("SELECT * FROM ai_usage WHERE invoice_id=?", (invoice_id,))
    pretty = json.dumps(json.loads(inv["extraction_json"]), indent=2) if inv["extraction_json"] else ""
    return page(request, "admin/invoice.html", inv=inv, lines=lines, usage=usage, pretty=pretty, tab="invoices",
                msg=request.query_params.get("msg"))


@app.post("/admin/release-held")
def admin_release_held(request: Request):
    """Queue every bill that arrived by email while the site was locked."""
    require_admin(request)
    n = db.x("UPDATE invoices SET status='queued' WHERE status='held'") and \
        db.q1("SELECT changes() n")["n"]
    db.x("UPDATE submissions SET notified_at=NULL WHERE notified_at IS NOT NULL AND notify_error='held'")
    return RedirectResponse("/admin/invoices?status=queued", 303)


@app.post("/admin/invoice/{invoice_id}/reprocess")
def admin_reprocess(request: Request, invoice_id: int):
    require_admin(request)
    db.x("UPDATE invoices SET status='queued', error=NULL, verified=0 WHERE id=?", (invoice_id,))
    return RedirectResponse(f"/admin/invoice/{invoice_id}?msg=Queued+for+reprocessing", 303)


@app.post("/admin/invoice/{invoice_id}/save")
def admin_save(request: Request, invoice_id: int, extraction_json: str = Form(...), verify: str = Form("")):
    """Correct an extraction by hand. Verified bills become reference examples for future extraction."""
    require_admin(request)
    try:
        data = json.loads(extraction_json)
    except Exception as e:
        return RedirectResponse(f"/admin/invoice/{invoice_id}?msg=Invalid+JSON:+{str(e)[:60]}", 303)
    extract.store(invoice_id, data)
    ok, diff = extract.reconcile(data)
    db.x("UPDATE invoices SET status='done', recon_ok=?, recon_diff=?, verified=?, confidence=? WHERE id=?",
         (int(ok), diff, 1 if verify else 0, 1.0 if verify else (0.97 if ok else 0.6), invoice_id))
    if verify and data.get("carrier"):
        extract.learn(data["carrier"], data, invoice_id)
    return RedirectResponse(f"/admin/invoice/{invoice_id}?msg=Saved", 303)


@app.get("/admin/visitors", response_class=HTMLResponse)
def admin_visitors(request: Request, hours: int = 168):
    require_admin(request)
    return page(request, "admin/visitors.html", rows=visitors.stories(hours), hours=hours, tab="visitors")


@app.post("/admin/digest-now")
def admin_digest_now(request: Request):
    require_admin(request)
    background(notify.daily_digest)
    return RedirectResponse("/admin/settings?saved=digest", 303)


@app.get("/setup/visitors")
def setup_visitors(token: str = "", hours: int = 168):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    return visitors.stories(hours)


@app.get("/admin/users", response_class=HTMLResponse)
def admin_users(request: Request):
    require_admin(request)
    rows = db.q("""
        SELECT e.email, u.name, u.company, u.created_at, u.last_login, u.password_hash IS NOT NULL AS registered,
               (SELECT COUNT(*) FROM invoices i WHERE i.owner_email=e.email) bills,
               (SELECT MAX(created_at) FROM invoices i WHERE i.owner_email=e.email) last_bill,
               (SELECT COUNT(*) FROM offer_requests r WHERE r.email=e.email) requests,
               (SELECT COUNT(*) FROM chat_messages c WHERE c.email=e.email AND role='user') questions
        FROM (SELECT email FROM users UNION SELECT DISTINCT owner_email FROM invoices WHERE owner_email IS NOT NULL) e
        LEFT JOIN users u ON u.email=e.email ORDER BY COALESCE(last_bill, u.created_at) DESC""")
    anon = db.q1("SELECT COUNT(DISTINCT session_id) n, COUNT(*) bills FROM invoices WHERE owner_email IS NULL")
    return page(request, "admin/users.html", rows=rows, anon=anon, tab="users")


@app.get("/admin/user", response_class=HTMLResponse)
def admin_user(request: Request, email: str):
    require_admin(request)
    user = db.q1("SELECT * FROM users WHERE email=?", (email.lower(),))
    return page(request, "dashboard.html", user=user, view_email=email.lower(), admin_view=True,
                msg=request.query_params.get("msg"), sales_email=db.get_setting("sales_email"))


@app.get("/admin/requests", response_class=HTMLResponse)
def admin_requests(request: Request):
    require_admin(request)
    rows = db.q("SELECT r.*, o.title offer_title, u.company FROM offer_requests r LEFT JOIN offers o ON o.id=r.offer_id"
                " LEFT JOIN users u ON u.email=r.email ORDER BY r.id DESC")
    return page(request, "admin/requests.html", rows=rows, tab="requests")


@app.post("/admin/requests/{rid}")
def admin_request_status(request: Request, rid: int, status: str = Form(...)):
    require_admin(request)
    db.x("UPDATE offer_requests SET status=? WHERE id=?", (status[:20], rid))
    return RedirectResponse("/admin/requests", 303)


@app.get("/admin/offers", response_class=HTMLResponse)
def admin_offers(request: Request):
    require_admin(request)
    return page(request, "admin/offers.html", rows=db.q("SELECT * FROM offers ORDER BY sort, id"), tab="offers")


@app.post("/admin/offers")
def admin_offers_save(request: Request, id: str = Form(""), title: str = Form(...), subtitle: str = Form(""),
                      details: str = Form(""), badge: str = Form(""), active: str = Form(""), sort: int = Form(0),
                      delete: str = Form("")):
    require_admin(request)
    oid = int(id) if id.isdigit() else None
    if oid and delete:
        db.x("DELETE FROM offers WHERE id=?", (oid,))
    elif oid:
        db.x("UPDATE offers SET title=?, subtitle=?, details=?, badge=?, active=?, sort=? WHERE id=?",
             (title, subtitle, details, badge, 1 if active else 0, sort, oid))
    else:
        db.x("INSERT INTO offers(title, subtitle, details, badge, active, sort, created_at) VALUES (?,?,?,?,?,?,?)",
             (title, subtitle, details, badge, 1 if active else 0, sort, time.time()))
    return RedirectResponse("/admin/offers", 303)


@app.get("/admin/knowledge", response_class=HTMLResponse)
def admin_knowledge(request: Request, carrier: str = ""):
    require_admin(request)
    return RedirectResponse("/admin/learning", 303)
    rows = db.q("SELECT * FROM knowledge WHERE (?='' OR carrier=?) ORDER BY carrier, kind, seen_count DESC",
                (carrier, carrier))
    return page(request, "admin/knowledge.html", rows=rows, carrier=carrier, tab="knowledge")


@app.post("/admin/knowledge/{kid}")
def admin_knowledge_edit(request: Request, kid: int, action: str = Form(...), value: str = Form("")):
    require_admin(request)
    if action == "delete":
        db.x("DELETE FROM knowledge WHERE id=?", (kid,))
    elif action == "toggle":
        db.x("UPDATE knowledge SET approved=1-approved WHERE id=?", (kid,))
    elif action == "save":
        db.x("UPDATE knowledge SET value=? WHERE id=?", (value[:300], kid))
    return RedirectResponse("/admin/learning", 303)


@app.post("/admin/knowledge")
def admin_knowledge_add(request: Request, carrier: str = Form(...), kind: str = Form(...), key: str = Form(...),
                        value: str = Form("")):
    require_admin(request)
    db.x("INSERT OR REPLACE INTO knowledge(carrier, kind, key, value, approved, updated_at) VALUES (?,?,?,?,1,?)",
         (carrier, kind, key[:160], value[:300], time.time()))
    return RedirectResponse("/admin/learning", 303)


EDITABLE = ["sales_email", "booking_url", "ai_provider", "claude_model", "chat_model", "local_ai_url",
            "local_ai_model", "mail_polling", "mail_user", "mail_address", "imap_host", "smtp_host",
            "alerts", "digest", "alert_email", "own_emails", "learning_provider", "openai_base_url", "openai_model",
            "openai_vision", "openai_vision_always", "openai_price_in", "openai_price_out", "site_lock", "site_password"]


@app.get("/admin/settings", response_class=HTMLResponse)
def admin_settings(request: Request):
    require_admin(request)
    hb = float(db.get_setting("worker_heartbeat") or 0)
    return page(request, "admin/settings.html", tab="settings",
                vals={k: db.get_setting(k) for k in EDITABLE},
                secrets_set={k: bool(db.get_setting(k)) for k in SECRET_SETTINGS},
                mail_status=json.loads(db.get_setting("mail_status") or "{}"),
                worker_age=time.time() - hb if hb else None, saved=request.query_params.get("saved"),
                last_test=db.get_setting("last_test_email"), deploy_log=_deploy_log(),
                last_ai_test=db.get_setting("last_ai_test"),
                held=db.q1("SELECT COUNT(*) n FROM invoices WHERE status='held'")["n"],
                branch=BRANCH, deploy_target=deploy_branch(), site_versions=SITE_VERSIONS,
                queue=db.q1("SELECT COUNT(*) n FROM invoices WHERE status IN ('queued','processing')")["n"])


def _deploy_log():
    p = DATA_DIR / "deploy.log"
    return "\n".join(p.read_text().splitlines()[-6:]) if p.exists() else ""


@app.post("/admin/settings")
async def admin_settings_save(request: Request):
    require_admin(request)
    form = await request.form()
    for k in EDITABLE:
        if k in form:
            db.set_setting(k, str(form[k]).strip())
    for k in SECRET_SETTINGS:
        v = str(form.get(k, "")).strip()
        if v:
            db.set_setting(k, v.replace(" ", "") if k == "mail_password" else v)
    return RedirectResponse("/admin/settings?saved=1", 303)


@app.post("/admin/site-version")
def admin_site_version(request: Request, branch: str = Form(...)):
    require_admin(request)
    set_deploy_branch(branch)
    return RedirectResponse("/admin/settings?saved=version", 303)


@app.get("/setup/backup")
def setup_backup(token: str = ""):
    """Diagnostics: a tar.gz of the database (consistent snapshot) and the bill PDFs, for moving servers."""
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    import io
    import sqlite3
    import tarfile
    import tempfile
    tmp = tempfile.NamedTemporaryFile(suffix=".sqlite", delete=False)
    tmp.close()
    src = sqlite3.connect(db.DB_PATH)
    dst = sqlite3.connect(tmp.name)
    src.backup(dst)
    dst.close()
    src.close()
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w:gz") as tar:
        tar.add(tmp.name, arcname="secondopinion.sqlite")
        if (DATA_DIR / "invoices").exists():
            tar.add(DATA_DIR / "invoices", arcname="invoices")
    os.unlink(tmp.name)
    from fastapi.responses import Response
    return Response(buf.getvalue(), media_type="application/gzip",
                    headers={"Content-Disposition": 'attachment; filename="secondopinion-backup.tar.gz"'})


@app.post("/setup/restore")
async def setup_restore(token: str = Form(...), file: UploadFile = File(...)):
    """Diagnostics: load a backup made by /setup/backup onto a fresh server (replaces the database)."""
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    import tarfile
    import tempfile
    data = await file.read()
    with tempfile.TemporaryDirectory() as td:
        with tarfile.open(fileobj=__import__("io").BytesIO(data), mode="r:gz") as tar:
            tar.extractall(td, filter="data")
        src = Path(td) / "secondopinion.sqlite"
        if not src.exists():
            raise HTTPException(400, "No database in the backup")
        for suffix in ("-wal", "-shm"):
            p = Path(str(db.DB_PATH) + suffix)
            if p.exists():
                p.unlink()
        __import__("shutil").copy(src, db.DB_PATH)
        inv = Path(td) / "invoices"
        n = 0
        if inv.exists():
            for f in inv.rglob("*.pdf"):
                dest = DATA_DIR / "invoices" / f.relative_to(inv)
                dest.parent.mkdir(parents=True, exist_ok=True)
                __import__("shutil").copy(f, dest)
                n += 1
    db.init()
    return {"ok": True, "pdfs": n, "invoices": db.q1("SELECT COUNT(*) n FROM invoices")["n"]}


@app.post("/setup/site-lock")
def setup_site_lock(token: str = Form(...), state: str = Form(...)):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    db.set_setting("site_lock", "on" if state == "on" else "off")
    return {"site_lock": db.get_setting("site_lock")}


@app.post("/setup/site-version")
def setup_site_version(token: str = Form(...), branch: str = Form(...)):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    set_deploy_branch(branch)
    return {"ok": True, "deploying": branch, "running": BRANCH}


@app.post("/admin/ai-test")
def admin_ai_test(request: Request):
    require_admin(request)
    from . import ai
    try:
        msg = ai.openai_test()
    except Exception as e:
        msg = f"failed: {type(e).__name__}: {e}"[:300]
    db.set_setting("last_ai_test", f"{time.strftime('%H:%M:%S')} {msg}")
    return RedirectResponse("/admin/settings?saved=aitest", 303)


@app.post("/admin/test-email")
def admin_test_email(request: Request):
    require_admin(request)
    try:
        mailer.send(SUPERADMIN_EMAIL, "SecondOpinion test email",
                    render("email/no_attachment.html", email=SUPERADMIN_EMAIL))
        msg = "sent OK"
    except Exception as e:
        msg = f"failed: {type(e).__name__}: {e}"[:200]
    db.set_setting("last_test_email", f"{time.strftime('%H:%M:%S')} {msg}")
    return RedirectResponse("/admin/settings?saved=test", 303)


@app.post("/admin/poll-now")
def admin_poll_now(request: Request):
    require_admin(request)
    background(mailer.poll)
    return RedirectResponse("/admin/settings?saved=poll", 303)


@app.exception_handler(StarletteHTTPException)
def http_error(request: Request, exc: StarletteHTTPException):
    if request.url.path.startswith("/api/"):
        return JSONResponse({"error": exc.detail}, exc.status_code)
    return page(request, "404.html", status_code=exc.status_code)


@app.get("/privacy", response_class=HTMLResponse)
def privacy(request: Request):
    return page(request, "privacy.html")


@app.get("/setup/inspect")
def setup_inspect(token: str = "", invoice_id: int = 0, pdf: int = 0):
    """Diagnostics for checking extraction quality (setup-token protected)."""
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    if not invoice_id:
        return db.q("SELECT i.id, s.source, s.notified_at, s.notify_error, owner_email, filename, status, error, carrier,"
                    " account_name, period_month, total_excl_gst,"
                    " line_count, recon_ok, recon_diff, pages, has_text, provider, model, duration_s, i.created_at"
                    " FROM invoices i LEFT JOIN submissions s ON s.id=i.submission_id ORDER BY i.id DESC LIMIT 30")
    inv = db.q1("SELECT * FROM invoices WHERE id=?", (invoice_id,))
    if not inv:
        raise HTTPException(404)
    if pdf:
        return FileResponse(inv["path"], media_type="application/pdf")
    _, text = extract.read_pdf(inv["path"])
    inv["extraction_json"] = json.loads(inv["extraction_json"]) if inv["extraction_json"] else None
    inv["pdf_text"] = text
    inv["usage"] = db.q("SELECT * FROM ai_usage WHERE invoice_id=?", (invoice_id,))
    return inv


@app.post("/setup/reprocess")
def setup_reprocess(token: str = Form(...), ids: str = Form(...), forget_value: str = Form(""),
                    resend: str = Form("")):
    """Re-queue bills for extraction; optionally drop learned facts with a given meaning (e.g. a retired category)."""
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    if forget_value:
        db.x("DELETE FROM knowledge WHERE value=?", (forget_value,))
    for i in [int(x) for x in ids.split(",") if x.strip().isdigit()]:
        if resend:
            db.x("UPDATE submissions SET notified_at=NULL, notify_error=NULL WHERE id=(SELECT submission_id FROM"
                 " invoices WHERE id=?)", (i,))
        else:
            db.x("UPDATE invoices SET status='queued', error=NULL WHERE id=?", (i,))
    return {"ok": True}


def activity(hours=24):
    since = time.time() - hours * 3600
    ev = db.q("SELECT kind, email, session_id, meta, created_at FROM events WHERE created_at>? ORDER BY id DESC",
              (since,))
    # Only views confirmed by the in-page beacon (a real browser that ran the page's script)
    views = [e for e in ev if e["kind"] == "page_view" and '"js": 1' in (e["meta"] or "")]
    metas = [json.loads(e["meta"] or "{}") for e in views]
    count = lambda k: sum(1 for e in ev if e["kind"] == k)
    tally = lambda key: sorted(
        ((v, sum(1 for m in metas if m.get(key) == v)) for v in {m.get(key) for m in metas if m.get(key)}),
        key=lambda t: -t[1])
    return {
        "hours": hours,
        "page_views": len(views),
        "visitors": len({e["session_id"] for e in views if e["session_id"]}),
        "uploads": count("upload"), "emails_in": count("email_received"), "registrations": count("registered"),
        "results_views": count("results_viewed"), "dashboard_views": count("dashboard_viewed"),
        "upload_intent": len({e["session_id"] for e in ev if e["kind"] == "upload_intent"}),
        "example_views": count("example_viewed"), "emails_captured": count("email_captured"),
        "ai_questions": count("chat"), "offer_requests": count("offer_request"), "errors": count("error"),
        "pages": tally("path")[:10], "referrers": tally("ref")[:10], "devices": tally("device"),
        "bills": db.q("SELECT status, COUNT(*) n FROM invoices WHERE created_at>? GROUP BY status", (since,)),
        "failed_bills": db.q("SELECT id, filename, owner_email, error FROM invoices WHERE created_at>?"
                             " AND status IN ('failed','rejected')", (since,)),
        "unsent_emails": db.q("SELECT id, email, notify_error FROM submissions WHERE created_at>?"
                              " AND notify_error IS NOT NULL", (since,)),
        "ai_cost_usd": db.q1("SELECT ROUND(COALESCE(SUM(cost_usd),0),2) v FROM ai_usage WHERE created_at>?",
                             (since,))["v"],
        "error_events": [json.loads(e["meta"] or "{}") | {"at": e["created_at"]} for e in ev if e["kind"] == "error"][:20],
        "feed": [{"kind": e["kind"], "email": e["email"], "meta": json.loads(e["meta"] or "{}"), "at": e["created_at"]}
                 for e in ev if e["kind"] != "page_view"][:60],
        "mail_status": json.loads(db.get_setting("mail_status") or "{}"),
        "worker_age_s": round(time.time() - float(db.get_setting("worker_heartbeat") or 0)),
    }


@app.get("/setup/stats")
def setup_stats(token: str = "", hours: float = 24):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    return activity(hours)


# ---------- learning & plan catalogue (superadmin) ----------

@app.get("/admin/learning", response_class=HTMLResponse)
def admin_learning(request: Request):
    require_admin(request)
    guides = db.q("SELECT * FROM format_guides ORDER BY carrier")
    audits = db.q("SELECT a.*, i.filename, i.carrier, i.account_name FROM audits a JOIN invoices i ON i.id=a.invoice_id"
                  " ORDER BY a.id DESC LIMIT 30")
    for a in audits:
        a["detail"] = json.loads(a["corrections"] or "{}")
    stats = db.q("SELECT i.carrier, COUNT(*) n, SUM(a.verdict='correct') correct, SUM(a.verdict='corrected') corrected,"
                 " SUM(a.applied) applied FROM audits a JOIN invoices i ON i.id=a.invoice_id GROUP BY i.carrier")
    parsers = db.q("SELECT * FROM parsers ORDER BY id DESC")
    for p in parsers:
        p["result"] = json.loads(p["last_result"] or "{}")
    suggestions = db.q("SELECT * FROM suggestions ORDER BY status='pending' DESC, id DESC LIMIT 60")
    knowledge = db.q("SELECT * FROM knowledge ORDER BY carrier, kind, seen_count DESC")
    return page(request, "admin/learning.html", tab="knowledge", guides=guides, audits=audits, stats=stats,
                parsers=parsers, suggestions=suggestions, rows=knowledge,
                learning_mode=db.get_setting("learning_mode"), min_bills=db.get_setting("parser_min_bills"),
                streak_needed=int(db.get_setting("parser_activate_streak") or 10))


@app.post("/admin/learning/guide")
def admin_guide_save(request: Request, carrier: str = Form(...), guide: str = Form(...)):
    require_admin(request)
    cur = db.q1("SELECT version FROM format_guides WHERE carrier=?", (carrier,))
    v = (cur["version"] + 1) if cur else 1
    db.x("INSERT INTO format_guides(carrier, guide, bills_seen, version, updated_at) VALUES (?,?,0,?,?)"
         " ON CONFLICT(carrier) DO UPDATE SET guide=excluded.guide, version=excluded.version,"
         " updated_at=excluded.updated_at", (carrier, guide, v, time.time()))
    db.x("INSERT INTO guide_history(carrier, version, guide, created_at) VALUES (?,?,?,?)",
         (carrier, v, guide, time.time()))
    return RedirectResponse("/admin/learning#guides", 303)


@app.post("/admin/learning/suggestion/{sug_id}")
def admin_suggestion(request: Request, sug_id: int, status: str = Form(...)):
    require_admin(request)
    db.x("UPDATE suggestions SET status=? WHERE id=?",
         (status if status in ("approved", "rejected", "pending") else "pending", sug_id))
    return RedirectResponse("/admin/learning#suggestions", 303)


@app.post("/admin/learning/parser/{pid}")
def admin_parser(request: Request, pid: int, action: str = Form(...)):
    require_admin(request)
    p = db.q1("SELECT * FROM parsers WHERE id=?", (pid,))
    if action == "activate" and p:
        db.x("UPDATE parsers SET status='retired' WHERE carrier=? AND status='active'", (p["carrier"],))
        db.x("UPDATE parsers SET status='active' WHERE id=?", (pid,))
    elif action == "retire":
        db.x("UPDATE parsers SET status='retired' WHERE id=?", (pid,))
    elif action == "backtest":
        background(learning.backtest, pid)
    return RedirectResponse("/admin/learning#parsers", 303)


@app.post("/admin/learning/mode")
def admin_learning_mode(request: Request, mode: str = Form(...)):
    require_admin(request)
    db.set_setting("learning_mode", "on" if mode == "on" else "off")
    return RedirectResponse("/admin/learning", 303)


@app.get("/admin/catalogue", response_class=HTMLResponse)
def admin_catalogue(request: Request):
    require_admin(request)
    rows = db.q("SELECT * FROM plan_catalogue WHERE active=1 ORDER BY provider, category, price_excl_gst")
    notes = {}
    for prov in catalogue.SITES:
        for cat in ("mobile", "broadband"):
            raw = db.get_setting(f"catalogue_note:{prov}:{cat}")
            notes[f"{prov}:{cat}"] = json.loads(raw) if raw else None
    brows = db.q("SELECT * FROM benefits WHERE active=1 ORDER BY provider, category, id")
    bnotes = {p: json.loads(db.get_setting(f"benefits_note:{p}") or "null") for p in catalogue.SITES}
    return page(request, "admin/catalogue.html", tab="catalogue", rows=rows, notes=notes, brows=brows, bnotes=bnotes,
                bstate=db.get_setting("benefits_requested"),
                state=db.get_setting("catalogue_requested"), last_result=db.get_setting("catalogue_last_result"),
                days=db.get_setting("catalogue_days"))


@app.post("/admin/benefits/refresh")
def admin_benefits_refresh(request: Request):
    require_admin(request)
    db.set_setting("benefits_requested", "1")
    background(_benefits_job)
    return RedirectResponse("/admin/catalogue#benefits", 303)


def _benefits_job():
    try:
        catalogue.refresh_all_benefits()
    finally:
        db.set_setting("benefits_requested", "0")


@app.post("/setup/benefits")
def setup_benefits(token: str = Form(...)):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    db.set_setting("benefits_requested", "1")
    background(_benefits_job)
    return {"ok": True}


@app.get("/setup/benefits")
def setup_benefits_list(token: str = ""):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    return {"rows": db.q("SELECT provider, category, title, value, detail, conditions, source_url FROM benefits"
                         " WHERE active=1 ORDER BY provider, category"),
            "state": db.get_setting("benefits_requested"), "last": db.get_setting("benefits_last_result")}


@app.post("/admin/catalogue/refresh")
def admin_catalogue_refresh(request: Request):
    require_admin(request)
    db.set_setting("catalogue_requested", "1")
    return RedirectResponse("/admin/catalogue", 303)


@app.post("/setup/catalogue")
def setup_catalogue(token: str = Form(...), provider: str = Form(""), category: str = Form("")):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    if provider:
        def job():
            try:
                for cat in ([category] if category else ["mobile", "broadband"]):
                    n = catalogue.refresh(provider, cat)
                    db.set_setting("catalogue_last_result", json.dumps({f"{provider}:{cat}": n}))
            except Exception as e:
                db.set_setting("catalogue_last_result", json.dumps({f"{provider}:{category}": f"error: {e}"[:200]}))
        background(job)
        return {"ok": True, "provider": provider}
    db.set_setting("catalogue_requested", "1")
    return {"ok": True}


@app.get("/setup/render-test")
def setup_render_test(token: str = "", url: str = "https://www.2degrees.nz/business/broadband/business-broadband"):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    return {"text": catalogue.render_page(url)[:3000]}


@app.get("/setup/learning")
def setup_learning(token: str = ""):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    return {
        "guides": db.q("SELECT carrier, version, bills_seen, updated_at, guide FROM format_guides"),
        "audits": db.q("SELECT invoice_id, verdict, applied, corrections FROM audits ORDER BY id DESC LIMIT 20"),
        "suggestions": db.q("SELECT * FROM suggestions ORDER BY id DESC LIMIT 20"),
        "notes": db.q("SELECT * FROM customer_notes"),
        "parsers": db.q("SELECT id, carrier, status, matches, mismatches, streak, last_result FROM parsers"),
        "catalogue": db.q("SELECT provider, category, COUNT(*) n, MAX(fetched_at) at FROM plan_catalogue"
                          " WHERE active=1 GROUP BY provider, category"),
        "catalogue_last_result": db.get_setting("catalogue_last_result"),
        "catalogue_state": db.get_setting("catalogue_requested"),
        "ai_by_purpose": db.q("SELECT purpose, COUNT(*) calls, ROUND(SUM(cost_usd),2) usd FROM ai_usage"
                              " GROUP BY purpose"),
    }


@app.get("/setup/catalogue-rows")
def setup_catalogue_rows(token: str = ""):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    return db.q("SELECT * FROM plan_catalogue WHERE active=1 ORDER BY provider, category, price_excl_gst")


# ---------- sales dossier ----------

@app.get("/admin/dossier.xlsx")
def admin_dossier_xlsx(request: Request, email: str):
    require_admin(request)
    from fastapi.responses import Response
    data, _ = dossier.build_xlsx(email.lower())
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{dossier.filename(email)}"'})


@app.post("/admin/dossier/send")
def admin_dossier_send(request: Request, email: str = Form(...)):
    require_admin(request)
    background(dossier.send, email.lower())
    return RedirectResponse(f"/admin/user?email={email}&msg=Dossier+is+on+its+way+to+{db.get_setting('sales_email')}", 303)


@app.get("/setup/dossier")
def setup_dossier(token: str = "", email: str = ""):
    """Diagnostics: the dossier's Services sheet as JSON."""
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    from openpyxl import load_workbook
    import io
    data, a = dossier.build_xlsx(email.lower())
    wb = load_workbook(io.BytesIO(data))
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


@app.post("/setup/dossier-send")
def setup_dossier_send(token: str = Form(...), email: str = Form(...)):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    background(dossier.send, email.lower())
    return {"ok": True, "to": db.get_setting("sales_email")}


def move_bills(from_email, to_email):
    """Re-home every bill (and derived notes/submissions) from one email address to another."""
    from_email, to_email = from_email.lower(), to_email.lower()
    with db.tx() as c:
        n = c.execute("UPDATE invoices SET owner_email=? WHERE owner_email=?", (to_email, from_email)).rowcount
        c.execute("UPDATE submissions SET email=? WHERE email=?", (to_email, from_email))
        c.execute("DELETE FROM customer_notes WHERE owner=?", (from_email,))
    ingest.dedupe(to_email)
    db.log_event("bills_moved", to_email, from_email=from_email, bills=n)
    return n


@app.post("/admin/move-bills")
def admin_move_bills(request: Request, from_email: str = Form(...), to_email: str = Form(...)):
    require_admin(request)
    n = move_bills(from_email, to_email)
    return RedirectResponse(f"/admin/user?email={to_email.lower()}&msg={n}+bills+moved+from+{from_email.lower()}", 303)


@app.post("/setup/move-bills")
def setup_move_bills(token: str = Form(...), from_email: str = Form(...), to_email: str = Form(...)):
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    return {"moved": move_bills(from_email, to_email)}


@app.post("/setup/delete-invoices")
def setup_delete_invoices(token: str = Form(...), ids: str = Form(...)):
    """Diagnostics: remove test bills (rows cascade; the PDF file stays on disk)."""
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    wanted = [int(x) for x in ids.split(",") if x.strip().isdigit()]
    with db.tx() as c:
        for i in wanted:
            c.execute("DELETE FROM invoices WHERE id=?", (i,))
    meter.refresh()
    return {"deleted": wanted}


@app.post("/setup/claim-invoice")
def setup_claim_invoice(token: str = Form(...), invoice_id: int = Form(...), email: str = Form(...)):
    """Diagnostics: attach one anonymous upload to an email address."""
    if not SETUP_TOKEN or token != SETUP_TOKEN:
        raise HTTPException(404)
    db.x("UPDATE invoices SET owner_email=? WHERE id=? AND owner_email IS NULL", (email.lower(), invoice_id))
    ingest.dedupe(email.lower())
    return {"ok": True}


@app.post("/setup/benefits-import")
async def setup_benefits_import(request: Request):
    """Diagnostics: load researched benefits rows (JSON list) without re-running the agent."""
    body = await request.json()
    if not SETUP_TOKEN or body.get("token") != SETUP_TOKEN:
        raise HTTPException(404)
    rows = body.get("rows") or []
    now = time.time()
    with db.tx() as c:
        for prov in {r["provider"] for r in rows}:
            c.execute("UPDATE benefits SET active=0 WHERE provider=?", (prov,))
        for r in rows:
            c.execute("INSERT INTO benefits(provider, category, title, detail, value, conditions, source_url, fetched_at)"
                      " VALUES (?,?,?,?,?,?,?,?)", (r["provider"], r["category"], r["title"], r.get("detail"),
                                                    r.get("value"), r.get("conditions"), r.get("source_url"), now))
    for prov in {r["provider"] for r in rows}:
        db.set_setting(f"benefits_note:{prov}", json.dumps({"at": now, "count": sum(1 for r in rows if r["provider"] == prov),
                                                          "notes": body.get("notes", "")}))
    return {"imported": len(rows)}
